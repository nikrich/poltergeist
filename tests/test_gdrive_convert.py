from __future__ import annotations

import io

import pytest

from ghostbrain.connectors.gdrive import convert, drive, tables
from tests.gdrive_fakes import (
    FakeDocs,
    FakeDrive,
    FakeSheets,
    drive_file,
    fake_services,
    http_error,
    install_fake_download,
)


def _docx(*paras: str) -> bytes:
    from docx import Document
    d = Document()
    for p in paras:
        d.add_paragraph(p)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _xlsx(sheets: dict[str, list[list]]) -> bytes:
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        for r in rows:
            ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_google_doc_uses_markdown_export():
    d = FakeDrive()
    d.exports["d1"] = b"# Title\n\nBody"
    res = convert.convert(fake_services(drive=d), drive_file("d1"))
    assert res.body == "# Title\n\nBody" and res.truncated is False


def test_big_google_doc_falls_back_to_docs_api():
    d = FakeDrive()
    d.too_big_exports.add("d1")
    docs = FakeDocs({"d1": {"body": {"content": [
        {"paragraph": {"elements": [{"textRun": {"content": "Huge doc\n"}}],
                       "paragraphStyle": {"namedStyleType": "HEADING_1"}}}]}}})
    res = convert.convert(fake_services(drive=d, docs=docs), drive_file("d1"))
    assert res.body == "# Huge doc"


def test_sheet_via_sheets_api_multi_tab():
    sh = FakeSheets({"s1": {"Budget": [["Item", "Cost"], ["Laptop", "1,200"]], "Empty": []}})
    res = convert.convert(fake_services(sheets=sh), drive_file("s1", mime=drive.GSHEET))
    assert res.body.startswith("## Budget\n\n| Item | Cost |")
    assert "Empty" not in res.body
    assert sh.batch_calls[0]["ranges"] == ["'Budget'!A1:AX5000", "'Empty'!A1:AX5000"]


def test_sheet_tab_names_are_quoted_in_ranges():
    sh = FakeSheets({"s1": {"Q3 'final'": [["a"]]}})
    convert.convert(fake_services(sheets=sh), drive_file("s1", mime=drive.GSHEET))
    assert sh.batch_calls[0]["ranges"] == ["'Q3 ''final'''!A1:AX5000"]


def test_sheet_truncation_from_grid_size():
    rows = [["v"] * tables.MAX_COLS for _ in range(tables.MAX_ROWS)]
    sh = FakeSheets({"s1": {"Log": rows}}, grid={"s1": {"Log": (9000, 80)}})
    res = convert.convert(fake_services(sheets=sh), drive_file("s1", mime=drive.GSHEET))
    assert res.truncated
    # A full 5,000x50 grid (even with 1-char cells) renders to ~1,010,000
    # chars on its own, already past MAX_BODY_CHARS (1,000,000). cap_body's
    # document-level cut lands inside the table body and replaces the
    # per-tab "_…truncated at N rows/columns_" notes with the generic
    # "_…truncated (document continues in Drive)_" marker. Both are correct
    # signals that the document was truncated; only the more specific one
    # survives here because the overall document cap bites first.
    assert "## Log" in res.body
    assert "_…truncated (document continues in Drive)_" in res.body


def test_sheet_tab_cap():
    sh = FakeSheets({"s1": {f"T{i}": [["v"]] for i in range(tables.MAX_TABS + 4)}})
    res = convert.convert(fake_services(sheets=sh), drive_file("s1", mime=drive.GSHEET))
    assert len(sh.batch_calls[0]["ranges"]) == tables.MAX_TABS
    assert res.body.endswith("_…4 more tabs_")


def test_xlsx_upload_formulas_become_values(monkeypatch, tmp_path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Calc"
    ws.append(["a", "b", "sum"])
    ws.append([1, 2, "=A2+B2"])
    p = tmp_path / "f.xlsx"
    wb.save(p)
    # openpyxl doesn't compute formulas; a saved-by-Excel file carries cached
    # values. Without a cache data_only gives None → the cell renders empty.
    d = FakeDrive()
    d.media["x1"] = p.read_bytes()
    install_fake_download(monkeypatch, d)
    res = convert.convert(fake_services(drive=d), drive_file("x1", mime=drive.XLSX, size=100))
    assert "## Calc" in res.body and "| 1 | 2 |" in res.body
    assert "=A2+B2" not in res.body


def test_xlsx_multi_tab(monkeypatch):
    d = FakeDrive()
    d.media["x1"] = _xlsx({"One": [["h"], ["r1"]], "Two": [["k"], ["r2"]]})
    install_fake_download(monkeypatch, d)
    res = convert.convert(fake_services(drive=d), drive_file("x1", mime=drive.XLSX, size=100))
    assert "## One" in res.body and "## Two" in res.body


def test_docx_upload(monkeypatch):
    d = FakeDrive()
    d.media["w1"] = _docx("Hello from Word")
    install_fake_download(monkeypatch, d)
    res = convert.convert(fake_services(drive=d), drive_file("w1", mime=drive.DOCX, size=100))
    assert "Hello from Word" in res.body


def test_empty_pdf_gets_placeholder(monkeypatch):
    d = FakeDrive()
    d.media["p1"] = b"%PDF-1.4 not really"
    install_fake_download(monkeypatch, d)
    res = convert.convert(fake_services(drive=d), drive_file("p1", mime=drive.PDF, size=100))
    assert res.body == convert.EMPTY_BODY


def test_too_large_skips_without_download(monkeypatch):
    called = []
    monkeypatch.setattr(drive, "download", lambda *a: called.append(a))
    with pytest.raises(convert.TooLarge):
        convert.convert(fake_services(), drive_file("p1", mime=drive.PDF, size=convert.MAX_DOWNLOAD_BYTES + 1))
    assert called == []


def test_cap_body_cuts_at_line_boundary():
    text = ("x" * 10 + "\n") * (convert.MAX_BODY_CHARS // 11 + 50)
    body, cut = convert.cap_body(text)
    assert cut and len(body) <= convert.MAX_BODY_CHARS + 60
    assert body.endswith("_…truncated (document continues in Drive)_")
    assert body.split("\n\n_…")[0].endswith("x" * 10)
    assert convert.cap_body("short") == ("short", False)


def test_other_http_errors_propagate():
    d = FakeDrive()
    d.fail_next = [http_error(404, "notFound")]
    from googleapiclient.errors import HttpError
    with pytest.raises(HttpError):
        convert.convert(fake_services(drive=d), drive_file("d1"))
